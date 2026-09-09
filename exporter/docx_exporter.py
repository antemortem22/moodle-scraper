"""Conversión estructural HTML -> Word editable; sin copiar el tema de Moodle."""
import base64
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote_to_bytes, urljoin

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

from .paths import unique_path


class DocxExporter:
    def __init__(self, downloader, warn=lambda message: None):
        self.downloader = downloader
        self.warn = warn

    def export(self, title: str, html: str, base_url: str, directory: Path, filename: str) -> Path:
        self.document = Document()
        self.document.styles['Normal'].font.name = 'Calibri'
        self.document.styles['Normal'].font.size = Pt(11)
        self.document.add_heading(title, level=1)
        self.base_url = base_url
        soup = BeautifulSoup(html, 'html.parser')
        for node in soup.select('script,style,nav,button,form,[hidden],[aria-hidden="true"]'):
            node.decompose()
        for node in list(soup.select('[style]')):
            if node.attrs is None:
                continue
            if re.search(r'(display\s*:\s*none|visibility\s*:\s*hidden)', node.get('style', ''), re.I):
                node.decompose()
        self.blocks(soup, self.document)
        path = unique_path(directory, filename)
        try:
            self.document.save(path)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        return path

    def image(self, node, paragraph):
        src = node.get('src') or node.get('data-src')
        if not src:
            return
        url = urljoin(self.base_url, src)
        try:
            if src.startswith('data:image/'):
                header, data = src.split(',', 1)
                data = base64.b64decode(data) if ';base64' in header else unquote_to_bytes(data)
            else:
                payload = self.downloader.fetch(url)
                if payload.is_html:
                    raise ValueError('El servidor devolvió HTML en lugar de una imagen.')
                data = payload.data
            # Pillow permite insertar también imágenes WebP/GIF como PNG en Word.
            stream = BytesIO()
            with Image.open(BytesIO(data)) as picture:
                width = min(6.0, picture.width / 96)
                picture.convert('RGBA').save(stream, format='PNG')
            stream.seek(0)
            paragraph.add_run().add_picture(stream, width=Inches(width))
        except Exception as exc:
            self.warn(f'Imagen no insertada: {exc}')
            paragraph.add_run(f"[Imagen: {node.get('alt') or 'sin descripción'}] ({url})")

    def inline(self, node, paragraph, bold=False, italic=False):
        if isinstance(node, NavigableString):
            text = re.sub(r'\s+', ' ', str(node))
            if text:
                run = paragraph.add_run(text)
                run.bold, run.italic = bold, italic
            return
        if not isinstance(node, Tag):
            return
        if node.name in ('p', 'div') and paragraph.text:
            paragraph.add_run().add_break()
        if node.name == 'img':
            self.image(node, paragraph)
            return
        if node.name == 'br':
            paragraph.add_run().add_break()
            return
        if node.name in ('iframe', 'object', 'embed', 'video', 'audio'):
            src = node.get('src') or node.get('data')
            source = node.find('source')
            src = src or (source.get('src') if source else None)
            paragraph.add_run(f"[{node.get('title') or 'Contenido multimedia'}]" + (f' ({urljoin(self.base_url, src)})' if src else ''))
            return
        if node.name == 'code':
            run = paragraph.add_run(node.get_text())
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            return
        style = node.get('style', '')
        bold = bold or node.name in ('strong','b') or bool(re.search(r'font-weight\s*:\s*(bold|[7-9]00)', style))
        italic = italic or node.name in ('em','i') or 'font-style: italic' in style
        for child in node.children:
            self.inline(child, paragraph, bold, italic)
        if node.name == 'a' and node.get('href'):
            paragraph.add_run(f" ({urljoin(self.base_url, node['href'])})")

    def list_numbering(self, ordered, start=1):
        numbering = self.document.part.numbering_part.element
        ids = [int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
        abstract_id = max(ids, default=-1) + 1
        abstract = OxmlElement('w:abstractNum')
        abstract.set(qn('w:abstractNumId'), str(abstract_id))
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), '0')
        for tag, value in [('start', str(start)), ('numFmt', 'decimal' if ordered else 'bullet'), ('lvlText', '%1.' if ordered else '•')]:
            el = OxmlElement('w:' + tag)
            el.set(qn('w:val'), value)
            lvl.append(el)
        abstract.append(lvl)
        numbering.append(abstract)
        return numbering.add_num(abstract_id).numId

    def blocks(self, parent, container, depth=0):
        paragraph = None
        for node in parent.children:
            if isinstance(node, NavigableString):
                if str(node).strip():
                    if paragraph is None:
                        paragraph = container.add_paragraph()
                    self.inline(node, paragraph)
                elif paragraph is not None:
                    self.inline(node, paragraph)
                continue
            if not isinstance(node, Tag):
                continue
            name = node.name
            if name in ('ul', 'ol'):
                start = int(node.get('start', '1')) if str(node.get('start', '1')).isdigit() else 1
                num_id = self.list_numbering(name == 'ol', start)
                for li in node.find_all('li', recursive=False):
                    p = container.add_paragraph(style='List Number' if name == 'ol' else 'List Bullet')
                    p.paragraph_format.left_indent = Inches(.25 * (depth + 1))
                    pr = p._p.get_or_add_pPr().get_or_add_numPr()
                    pr.get_or_add_ilvl().val = 0
                    pr.get_or_add_numId().val = num_id
                    # Flush groups around nested lists to preserve their exact order.
                    for child in li.children:
                        if isinstance(child, Tag) and child.name in ('ul', 'ol', 'table'):
                            wrapper = BeautifulSoup(str(child), 'html.parser')
                            self.blocks(wrapper, container, depth + 1)
                            p = None
                        else:
                            if p is None:
                                p = container.add_paragraph()
                            self.inline(child, p)
                paragraph = None
            elif name == 'table':
                self.table(node, container)
                paragraph = None
            elif name == 'pre':
                p = container.add_paragraph()
                p.paragraph_format.left_indent = Inches(.2)
                run = p.add_run(node.get_text())
                run.font.name, run.font.size = 'Consolas', Pt(10)
                paragraph = None
            elif re.fullmatch(r'h[1-6]', name):
                p = container.add_paragraph(style=f'Heading {min(int(name[1])+1, 6)}')
                for child in node.children:
                    self.inline(child, p)
                paragraph = None
            elif name in ('p', 'blockquote'):
                p = container.add_paragraph(style='Quote' if name == 'blockquote' else 'Normal')
                if name == 'blockquote' and node.find(['p', 'table', 'ul', 'ol']):
                    before = len(container.paragraphs)
                    self.blocks(node, container, depth)
                    for child_paragraph in container.paragraphs[before:]:
                        if child_paragraph.style.name == 'Normal':
                            child_paragraph.style = 'Quote'
                else:
                    for child in node.children:
                        self.inline(child, p)
                paragraph = None
            elif name in ('div','section','article','main','body','html','figure','figcaption','details','summary'):
                self.blocks(node, container, depth)
                paragraph = None
            else:
                if paragraph is None:
                    paragraph = container.add_paragraph()
                self.inline(node, paragraph)

    def table(self, node, container):
        caption = node.find('caption', recursive=False)
        if caption:
            container.add_paragraph(caption.get_text(' ', strip=True), style='Caption')
        rows = [r for r in node.find_all('tr') if r.find_parent('table') is node]
        if not rows:
            return
        occupied, placements = set(), []
        columns = 0
        for row_index, row in enumerate(rows):
            col = 0
            for cell in row.find_all(['td','th'], recursive=False):
                while (row_index,col) in occupied:
                    col += 1
                colspan = cell.get('colspan','1')
                rowspan = cell.get('rowspan','1')
                colspan = max(1,min(100,int(colspan))) if str(colspan).isdigit() else 1
                rowspan = max(1,min(len(rows)-row_index,int(rowspan))) if str(rowspan).isdigit() else 1
                placements.append((row_index,col,rowspan,colspan,cell))
                occupied.update((r,c) for r in range(row_index,row_index+rowspan) for c in range(col,col+colspan))
                col += colspan
                columns = max(columns,col)
        if not columns:
            return
        table = container.add_table(rows=len(rows), cols=columns)
        table.style = 'Table Grid'
        for r,c,rs,cs,node in placements:
            cell = table.cell(r,c)
            if rs > 1 or cs > 1:
                cell = cell.merge(table.cell(r+rs-1,c+cs-1))
            self.blocks(node, cell)
            if len(cell.paragraphs) > 1 and not cell.paragraphs[0].text:
                empty = cell.paragraphs[0]._p
                empty.getparent().remove(empty)
            if node.name == 'th':
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
