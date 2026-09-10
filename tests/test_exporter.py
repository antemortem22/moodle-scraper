import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from PIL import Image

from exporter.docx_exporter import DocxExporter
from exporter.downloader import Downloader, Payload
from exporter.paths import safe_name, unique_path
from exporter.service import academic_content, export_sections
from models import Activity, Course, Section


class Response:
    def __init__(self, url, data, headers=None, status=200):
        self.url, self.data, self.headers, self.status = url, data, headers or {}, status
        self.ok = 200 <= status < 300
    def body(self):
        return self.data
    def dispose(self):
        pass


class Request:
    def __init__(self, responses):
        self.responses, self.calls = responses, []
    def get(self, url, **kwargs):
        self.calls.append(url)
        return self.responses[url]


class ExporterTests(unittest.TestCase):
    def test_windows_names_and_collisions(self):
        self.assertEqual(safe_name('CON.txt'), '_CON.txt')
        self.assertNotIn(':', safe_name('Tema: 1?'))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first = unique_path(root, 'Texto.docx')
            first.write_bytes(b'original')
            second = unique_path(root, 'Texto.docx')
            self.assertNotEqual(first, second)
            self.assertEqual(first.read_bytes(), b'original')
            self.assertEqual(unique_path(root, '../archivo.pdf').parent, root)

    def test_redirect_original_filename_and_no_external_url_fetch(self):
        requests = Request({
            'https://m/mod/resource/view.php?id=1': Response('https://m/mod/resource/view.php?id=1', b'', {'location':'/file'},302),
            'https://m/file': Response('https://m/file', b'%PDF-original', {'content-type':'application/pdf','content-disposition':"attachment; filename*=UTF-8''Gu%C3%ADa.pdf"}),
            'https://m/mod/url/view.php?id=2': Response('', b'', {'location':'https://external.test/site'},302),
        })
        downloader = Downloader(requests)
        with tempfile.TemporaryDirectory() as directory:
            path = downloader.download('https://m/mod/resource/view.php?id=1', Path(directory), '02 - ', 'fallback')
            self.assertEqual(path.name, '02 - Guía.pdf')
            self.assertEqual(path.read_bytes(), b'%PDF-original')
        self.assertEqual(downloader.external_url('https://m/mod/url/view.php?id=2'), 'https://external.test/site')
        self.assertNotIn('https://external.test/site', requests.calls)
        payload = Payload('https://m/file', {'content-disposition':'attachment; filename="GuÃ­a.pdf"'}, b'')
        self.assertEqual(downloader.filename(payload, 'fallback'), 'Guía.pdf')

    def test_academic_content_excludes_submission_ui(self):
        payload = Payload('https://m/assign', {'content-type':'text/html'}, b'''<div id="intro"><p>Consigna</p><div id="assign_files_tree12"><a href="/pluginfile.php/1/mod_assign/introattachment/0/task.pdf">task.pdf</a></div></div><div role="main">Calificaciones y entregas</div>''')
        html, files = academic_content(payload, 'assign')
        self.assertIn('Consigna', html)
        self.assertNotIn('Calificaciones', html)
        self.assertNotIn('assign_files_tree', html)
        self.assertEqual(len(files), 1)

    def test_word_structure_images_and_failed_image(self):
        picture = BytesIO()
        Image.new('RGB', (80, 40), 'red').save(picture, format='PNG')
        requests = Request({'https://m/image.png': Response('https://m/image.png', picture.getvalue(), {'content-type':'image/png'})})
        warnings = []
        html = '''<h1>Subtítulo</h1><p>Texto <strong>negrita</strong> <em>cursiva</em>.</p>
        <ul><li>Uno</li><li>Dos<ul><li>Anidado</li></ul></li></ul><ol><li>Primero</li></ol>
        <table><tr><th colspan="2">Cabecera</th></tr><tr><td>A</td><td>B</td></tr></table>
        <blockquote>Cita</blockquote><pre>print(1)\n  indentado</pre>
        <p><img src="/image.png"><a href="/link">Enlace</a></p>
        <iframe src="https://external.test/interactive" title="Interactivo"></iframe>
        <p><img src="/missing.png" alt="Imagen faltante"></p>'''
        with tempfile.TemporaryDirectory() as directory:
            path = DocxExporter(Downloader(requests), warnings.append).export('Título',html,'https://m/page',Path(directory),'01 - Texto.docx')
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].style.name, 'Heading 1')
            self.assertEqual(doc.paragraphs[1].style.name, 'Heading 2')
            self.assertTrue(any(r.bold for p in doc.paragraphs for r in p.runs))
            self.assertTrue(any(r.italic for p in doc.paragraphs for r in p.runs))
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.inline_shapes), 1)
            text = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('https://m/link', text)
            self.assertIn('https://external.test/interactive', text)
            self.assertIn('Imagen faltante', text)
            self.assertIn('print(1)\n  indentado', text)
            with ZipFile(path) as archive:
                self.assertIn(b'numPr', archive.read('word/document.xml'))
                self.assertIn(b'gridSpan', archive.read('word/document.xml'))
            self.assertEqual(len(warnings),1)
            self.assertNotIn('https://external.test/interactive', requests.calls)

    def test_export_continues_without_generating_indexes(self):
        course = Course(1,'Curso','https://m/course/view.php?id=1')
        section = Section('Semana',1,course.url+'&section=1')
        activities = [
            Activity(1,'Error','https://m/bad','resource',section,1),
            Activity(2,'Archivo','https://m/file','resource',section,2),
            Activity(3,'No soportado','https://m/quiz','quiz',section,3),
            Activity(4,'Restringida','https://m/restricted','page',section,4,False),
            Activity(5,'Texto',None,'label',section,5,content_html='<p>Texto académico</p>'),
        ]
        request = Request({'https://m/bad':Response('https://m/bad',b'',status=403),
                           'https://m/file':Response('https://m/file',b'original',{'content-disposition':'attachment; filename="file.zip"'})})
        page = SimpleNamespace(context=SimpleNamespace(request=request))
        with tempfile.TemporaryDirectory() as directory, patch('exporter.service.get_activities',return_value=activities) as discover:
            log_path = Path(directory)/'logs'/'export.log'
            report = export_sections(page,course,[section],output_root=directory,log_path=log_path)
            discover.assert_called_once_with(page,course,[section])
            self.assertEqual(len(report.errors),1)
            self.assertEqual([p.name for p in report.files], ['02 - file.zip', '05 - Texto.docx'])
            self.assertEqual(list(report.directory.rglob('00 - *.docx')), [])
            self.assertNotIn('https://m/restricted',request.calls)
            self.assertNotIn('https://m/quiz',request.calls)
            self.assertEqual(list(report.directory.rglob('*.log')), [])
            self.assertIn('[CURSO] Curso', log_path.read_text(encoding='utf-8'))
            self.assertIn('[ERROR]', log_path.read_text(encoding='utf-8'))

    def test_assign_and_folder_continue_after_attachment_error(self):
        course = Course(1,'Curso','https://m/course/view.php?id=1')
        section = Section('Semana',1,course.url+'&section=1')
        intro = b'<div id="intro"><p>Consigna</p><a href="/pluginfile.php/a">A</a><a href="/pluginfile.php/b">B</a></div><div role="main"></div>'
        folder = b'<div role="main"><div class="foldertree"><a href="/pluginfile.php/c">C</a><a href="/pluginfile.php/d">D</a></div></div>'
        request = Request({
            'https://m/assign':Response('https://m/assign',intro,{'content-type':'text/html'}),
            'https://m/folder':Response('https://m/folder',folder,{'content-type':'text/html'}),
            'https://m/pluginfile.php/a':Response('',b'',status=403),
            **{f'https://m/pluginfile.php/{key}':Response(f'https://m/pluginfile.php/{key}',key.encode(),{'content-disposition':'attachment; filename="original.xlsx"'}) for key in ('b','c','d')},
        })
        items = [Activity(1,'Consigna','https://m/assign','assign',section,1),Activity(2,'Carpeta','https://m/folder','folder',section,2)]
        page = SimpleNamespace(context=SimpleNamespace(request=request))
        with tempfile.TemporaryDirectory() as directory, patch('exporter.service.get_activities',return_value=items):
            report = export_sections(page,course,[section],output_root=directory,log_path=Path(directory)/'logs'/'export.log')
            self.assertEqual(len(report.errors),1)
            self.assertEqual(len(report.files),4)
            self.assertEqual([p.read_bytes() for p in report.files if p.suffix=='.xlsx'],[b'b',b'c',b'd'])


if __name__ == '__main__':
    unittest.main()
